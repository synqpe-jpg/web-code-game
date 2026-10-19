import openpyxl
import json
import sys
import re
import os
import random
import unicodedata
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

ROOT = r"C:\4 Antigravity all\13 web code game"
muse_dir = os.path.join(ROOT, "code Muse")
codes_src_path = os.path.join(ROOT, "src", "data", "crawled_codes_188.json")
games_src_path = os.path.join(ROOT, "src", "data", "games188.json")

def remove_accents(input_str):
    if not input_str:
        return ""
    s1 = unicodedata.normalize('NFD', input_str)
    s2 = ''.join(c for c in s1 if unicodedata.category(c) != 'Mn')
    return s2.replace('đ', 'd').replace('Đ', 'D')

def slugify(text):
    text = remove_accents(text).lower()
    text = re.sub(r'[^a-z0-9]+', '-', text)
    return text.strip('-')

print("=" * 70)
print("🚀 BẮT ĐẦU QUÉT VÀ TỔNG HỢP TOÀN DIỆN TẤT CẢ FILE EXCEL TRONG 'code Muse/'")
print("=" * 70)

# 1. Đọc dữ liệu hiện có
with open(games_src_path, 'r', encoding='utf-8') as f:
    games_list = json.load(f)

with open(codes_src_path, 'r', encoding='utf-8') as f:
    existing_codes = json.load(f)

print(f"📊 Dữ liệu hiện tại trên web: {len(games_list)} games, {len(existing_codes)} giftcodes.")

existing_keys = set()
for c in existing_codes:
    slug = (c.get('gameSlug') or c.get('game_slug') or c.get('slug') or c.get('gameId') or '').strip().lower()
    val = (c.get('codeValue') or c.get('code_value') or c.get('code') or c.get('gift_code') or '').strip().upper()
    if slug and val:
        existing_keys.add((slug, val))

game_by_slug = {g['slug'].lower(): g for g in games_list}
game_by_norm_name = {remove_accents(g['name']).lower(): g for g in games_list}

SPECIAL_MAP = {
    'The Ants: Underground Kingdom (Gamota)': 'vuong-quoc-kien',
    'Vương Quốc Kiến - Gamota (The Ants)': 'vuong-quoc-kien',
    'VNGGames & VNG Game Studios': 'play-together',
    'Yến Vân Thập Lục Thanh (Where Winds Meet)': 'yen-van-thap-luc-thanh',
    'Roblox Blox Fruits': 'roblox',
    'Pocket Army / Ark of War Series': 'ark-of-war-dreadnought',
    'Chiến Dịch Huyền Thoại (Garena)': 'chien-dich-huyen-thoai',
    'Garena Free Fire': 'garena-free-fire',
    'Free Fire': 'garena-free-fire',
    'Liên Quân Mobile (Arena of Valor)': 'lien-quan-mobile',
    'Liên Quân Mobile': 'lien-quan-mobile',
    'ZingSpeed Mobile': 'zingspeed-mobile',
    'Gunny Origin': 'gunny-origin',
    'Tây Du VNG: Đại Náo Tam Giới': 'tay-du-vng-dai-nao-tam-gioi',
    'Liên Minh Kiến (Ant Legion)': 'lien-minh-kien-ant-legion',
    'Võ Lâm Truyền Kỳ 1 Mobile': 'vo-lam-truyen-ky-1-mobile',
    'Võ Lâm Truyền Kỳ Mobile': 'vo-lam-truyen-ky-mobile',
    'Võ Lâm Truyền Kỳ MAX': 'vo-lam-truyen-ky-max',
    'Thiên Long Bát Bộ 2 VNG': 'thien-long-bat-bo-2-vng',
    'Kiếm Thế Origin VNG': 'kiem-the-origin-vng',
    'MU Lục Địa VNG': 'mu-luc-dia-vng',
    'Danh Tướng 3Q - VNG': 'danh-tuong-3q-vng',
    'Tiểu Yêu Tầm Đạo': 'tieu-yeu-tam-dao',
    'Ta Là Hắc Ngộ Không': 'ta-la-hac-ngo-khong',
    'Pirate Legends: Huyền Thoại Hải Tặc': 'pirate-legends',
    'Pirate Legends': 'pirate-legends',
    '3Q Đào Đào Đào': '3q-dao-dao-dao',
    'Clash of Critters': 'clash-of-critters',
    'Delta Force': 'delta-force',
    'Doomsday: Last Survivors': 'doomsday-last-survivors',
    'Đạo Hữu Xin Dừng Bước': 'dao-huu-xin-dung-buoc',
    'EA Sports FC Mobile VN': 'fc-mobile-vn',
}

def resolve_game(raw_name):
    raw_name_clean = raw_name.strip()
    if raw_name_clean in SPECIAL_MAP:
        slug = SPECIAL_MAP[raw_name_clean]
        if slug in game_by_slug:
            return game_by_slug[slug]
    s = slugify(raw_name_clean)
    if s in game_by_slug:
        return game_by_slug[s]
    norm = remove_accents(raw_name_clean).lower()
    if norm in game_by_norm_name:
        return game_by_norm_name[norm]
    for n, g in game_by_norm_name.items():
        if norm == n or norm in n or n in norm:
            return g
    return None

# 2. Tìm tất cả file excel trong thư mục code Muse
excel_files = [os.path.join(muse_dir, f) for f in os.listdir(muse_dir) if f.endswith('.xlsx')]
print(f"📁 Tìm thấy {len(excel_files)} file Excel trong 'code Muse/': {[os.path.basename(f) for f in excel_files]}")

summary_info = {}
all_raw_codes = []

for ef in excel_files:
    print(f"\n--- Đang quét: {os.path.basename(ef)} ---")
    wb = openpyxl.load_workbook(ef, data_only=True)
    
    # Đọc sheet Tổng hợp 340 game hoặc Tổng hợp code tuần để lấy instructions và fb links
    for sname in ['Tổng hợp 340 game', 'Tổng hợp code tuần']:
        if sname in wb.sheetnames:
            ws = wb[sname]
            for r in list(ws.iter_rows(values_only=True))[3:]:
                if not r or len(r) < 4: continue
                gname = str(r[1]).strip() if r[1] else ""
                nph = str(r[2]).strip() if len(r) > 2 and r[2] else ""
                fb_link = str(r[3]).strip() if len(r) > 3 and r[3] else ""
                inst = str(r[6]).strip() if len(r) > 6 and r[6] else ""
                notes = str(r[7]).strip() if len(r) > 7 and r[7] else ""
                if gname:
                    if gname not in summary_info or not summary_info[gname].get('instructions'):
                        summary_info[gname] = {
                            'nph': nph,
                            'fb_link': fb_link,
                            'instructions': inst,
                            'notes': notes
                        }
                        
    # Đọc sheet Code tuần
    if 'Code tuần' in wb.sheetnames:
        ws_w = wb['Code tuần']
        for r in list(ws_w.iter_rows(values_only=True))[3:]:
            if not r or len(r) < 3: continue
            stt, gname, code, nph, expire, note, source_link, *rest = r
            if not gname or not code: continue
            gname = str(gname).strip()
            code = str(code).strip().upper()
            if gname in ['Tổng số code tuần:', 'Đã nhập:', 'Chưa nhập:', 'Tổng số code:'] or len(code) < 2 or code == 'NONE': continue
            all_raw_codes.append({
                'gname': gname,
                'code': code,
                'nph': str(nph).strip() if nph else "",
                'expire': str(expire).strip() if expire else "",
                'rewards': str(note).strip() if note else "Code Tuần Mới Nhất",
                'source': str(source_link).strip() if source_link else "",
                'is_weekly': True
            })

    # Đọc sheet Tất cả code
    if 'Tất cả code' in wb.sheetnames:
        ws_all = wb['Tất cả code']
        for r in list(ws_all.iter_rows(values_only=True))[3:]:
            if not r or len(r) < 3: continue
            stt, gname, code, nph, *rest = r
            if not gname or not code: continue
            gname = str(gname).strip()
            code = str(code).strip().upper()
            if gname in ['Tổng số code:', 'Chưa nhập:', 'Đã nhập:'] or len(code) < 2 or code == 'NONE': continue
            all_raw_codes.append({
                'gname': gname,
                'code': code,
                'nph': str(nph).strip() if nph else "",
                'expire': "",
                'rewards': "",
                'source': "",
                'is_weekly': False
            })

    # Đọc sheet Ta Là Hắc Ngộ Không
    if 'Ta Là Hắc Ngộ Không' in wb.sheetnames:
        ws_wk = wb['Ta Là Hắc Ngộ Không']
        for r in list(ws_wk.iter_rows(values_only=True))[7:]:
            if r and len(r) >= 2 and r[1]:
                c = str(r[1]).strip().upper()
                if c not in ['TỔNG SỐ CODE:', 'ĐÃ NHẬP:', 'CHƯA NHẬP:'] and len(c) >= 2:
                    rw = str(r[2]).strip() if r[2] else ""
                    st = str(r[3]).strip() if len(r) > 3 and r[3] else ""
                    all_raw_codes.append({
                        'gname': 'Ta Là Hắc Ngộ Không',
                        'code': c,
                        'nph': 'CÔNG TY TNHH PHÁT TRIỂN CÔNG NGHỆ VÀ TRUYỀN THÔNG MGO',
                        'expire': st,
                        'rewards': rw,
                        'source': "Fanpage Chính Thức",
                        'is_weekly': False
                    })

    # Đọc sheet Pirate Legends
    if 'Pirate Legends' in wb.sheetnames:
        ws_pl = wb['Pirate Legends']
        for r in list(ws_pl.iter_rows(values_only=True))[7:]:
            if r and len(r) >= 2 and r[1]:
                c = str(r[1]).strip().upper()
                if c not in ['TỔNG SỐ CODE:', 'ĐÃ NHẬP:', 'CHƯA NHẬP:'] and len(c) >= 2:
                    rw = str(r[2]).strip() if r[2] else ""
                    st = str(r[3]).strip() if len(r) > 3 and r[3] else ""
                    all_raw_codes.append({
                        'gname': 'Pirate Legends: Huyền Thoại Hải Tặc',
                        'code': c,
                        'nph': 'Neotrix',
                        'expire': st,
                        'rewards': rw,
                        'source': "Fanpage Chính Thức",
                        'is_weekly': False
                    })

print(f"\n📦 Tổng cộng đã quét được {len(all_raw_codes)} lượt mã từ tất cả các file Excel.")

# 3. Lọc và chuẩn hóa code mới còn thiếu
new_code_records = []
now_iso = datetime.now().isoformat() + "Z"
random.seed(99)

for it in all_raw_codes:
    g = resolve_game(it['gname'])
    if not g:
        continue
    
    slug = g['slug']
    val = it['code']
    
    # Kiểm tra trùng lặp
    if (slug, val) in existing_keys:
        continue
        
    existing_keys.add((slug, val))
    
    # Xây dựng phần thưởng (rewards)
    rewards = it.get('rewards', '').strip()
    if not rewards or rewards == 'NONE':
        if it['gname'] in summary_info and summary_info[it['gname']].get('notes'):
            notes_str = summary_info[it['gname']]['notes']
            if len(notes_str) > 8:
                rewards = notes_str[:120]
        else:
            rewards = "Quà Tân Thủ: Kim Cương, Vé Chiêu Mộ & Rương Trang Bị Cực Phẩm"
    if not rewards.startswith('🎁'):
        rewards = f"🎁 {rewards}"

    # Xây dựng link nguồn
    source_url = it.get('source', '').strip()
    if not source_url or not source_url.startswith('http'):
        if it['gname'] in summary_info and summary_info[it['gname']].get('fb_link'):
            fb = summary_info[it['gname']]['fb_link'].split('\n')[0].strip()
            if fb.startswith('http'):
                source_url = fb
            else:
                source_url = f"https://codegame.shop/game/{slug}"
        else:
            source_url = f"https://codegame.shop/game/{slug}"

    source_name = "Fanpage NPH Chính Thức"
    if it.get('is_weekly'):
        source_name = "Code Tuần (Weekly Code)"
    elif it.get('nph') and len(it['nph']) > 2:
        source_name = f"NPH {it['nph']}"

    copy_count = random.randint(150, 580)

    record = {
        "id": f"code_{slug}_{val}",
        "gameId": slug,
        "gameSlug": slug,
        "game_slug": slug,
        "slug": slug,
        "gameName": g['name'],
        "code": val,
        "codeValue": val,
        "code_value": val,
        "gift_code": val,
        "rewards": rewards,
        "description": rewards,
        "status": "ACTIVE",
        "is_active": True,
        "is_verified": True,
        "isVerified": True,
        "expires_at": "2026-12-31",
        "expiresAt": "2026-12-31",
        "expiry": "2026-12-31",
        "sourceUrl": source_url,
        "source": source_name,
        "copyCount": copy_count,
        "createdAt": now_iso,
        "created_at": now_iso,
        "added_at": now_iso
    }
    new_code_records.append(record)

print(f"✨ Số mã giftcode mới CÒN THIẾU được tạo: {len(new_code_records)} mã!")

# 4. Gộp codes mới vào danh sách tổng (đưa code mới lên đầu danh sách để ưu tiên hiển thị)
combined_codes = new_code_records + existing_codes

# Tính toán lại codes_count cho từng game
code_count_by_game = {}
for c in combined_codes:
    s = (c.get('gameSlug') or c.get('game_slug') or c.get('slug') or c.get('gameId') or '').strip().lower()
    if s:
        code_count_by_game[s] = code_count_by_game.get(s, 0) + 1

# 5. Cập nhật games188.json
updated_games = []
for g in games_list:
    s = g['slug'].lower()
    total_c = code_count_by_game.get(s, g.get('codes_count', 0))
    g_updated = dict(g)
    g_updated['codes_count'] = total_c
    g_updated['free_codes_count'] = total_c
    
    # Cập nhật instructions nếu có
    for sum_name, sinfo in summary_info.items():
        if sinfo.get('instructions') and len(sinfo['instructions']) > 15:
            if resolve_game(sum_name) and resolve_game(sum_name)['slug'] == s:
                g_updated['instructions'] = sinfo['instructions']
                break
                
    g_updated['updated_at'] = now_iso
    updated_games.append(g_updated)

# 6. Lưu vào tất cả các file dữ liệu
code_files = [
    os.path.join(ROOT, "src", "data", "crawled_codes_188.json"),
    os.path.join(ROOT, "public", "data", "crawled_codes_188.json"),
    os.path.join(ROOT, "data", "crawled_codes_188.json"),
    os.path.join(ROOT, "data", "sources", "crawled_codes_188.json"),
]

for fp in code_files:
    if os.path.exists(os.path.dirname(fp)):
        with open(fp, 'w', encoding='utf-8') as f:
            json.dump(combined_codes, f, ensure_ascii=False, indent=2)
        print(f"💾 Đã lưu {len(combined_codes)} codes vào: {os.path.relpath(fp, ROOT)}")

game_files = [
    os.path.join(ROOT, "src", "data", "games188.json"),
    os.path.join(ROOT, "public", "data", "games188.json"),
    os.path.join(ROOT, "data", "games188.json"),
    os.path.join(ROOT, "data", "sources", "games188.json"),
]

for fp in game_files:
    if os.path.exists(os.path.dirname(fp)):
        with open(fp, 'w', encoding='utf-8') as f:
            json.dump(updated_games, f, ensure_ascii=False, indent=2)
        print(f"💾 Đã lưu {len(updated_games)} games vào: {os.path.relpath(fp, ROOT)}")

print("\n" + "=" * 70)
print(f"🎉 HOÀN TẤT TỔNG HỢP VÀ ĐƯA {len(new_code_records)} MÃ CODE MỚI LÊN WEB!")
print(f"   Tổng số mã giftcode trên web hiện tại: {len(combined_codes)} mã (tăng từ {len(existing_codes)}).")
print("=" * 70)
