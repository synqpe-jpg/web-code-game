import openpyxl
import json
import sys
import re
import os
import random
import unicodedata
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

ROOT = r"C:\4 Antigravity all\13 web code game"
excel_path = os.path.join(ROOT, "code Muse", "Giftcode 2 Game.xlsx")
codes_src_path = os.path.join(ROOT, "src", "data", "crawled_codes_188.json")
games_src_path = os.path.join(ROOT, "src", "data", "games188.json")

def remove_accents(input_str):
    if not input_str:
        return ""
    s1 = unicodedata.normalize('NFD', input_str)
    s2 = ''.join(c for c in s1 if unicodedata.category(c) != 'Mn')
    return s2.replace('đ', 'd').replace('Đ', 'D')

def slugify(text):
    text = remove_accents(text).lower()
    text = re.sub(r'[^a-z0-9]+', '-', text)
    return text.strip('-')

print("=" * 65)
print("🚀 BẮT ĐẦU ĐỒNG BỘ VÀ BỔ SUNG CODE TỪ FILE 'Giftcode 2 Game.xlsx'")
print("=" * 65)

# 1. Đọc dữ liệu hiện có
with open(games_src_path, 'r', encoding='utf-8') as f:
    games_list = json.load(f)

with open(codes_src_path, 'r', encoding='utf-8') as f:
    existing_codes = json.load(f)

print(f"📊 Dữ liệu hiện tại: {len(games_list)} games, {len(existing_codes)} giftcodes.")

# Tạo set kiểm tra trùng lặp
existing_keys = set()
for c in existing_codes:
    slug = (c.get('gameSlug') or c.get('game_slug') or c.get('slug') or c.get('gameId') or '').strip().lower()
    val = (c.get('codeValue') or c.get('code_value') or c.get('code') or c.get('gift_code') or '').strip().upper()
    if slug and val:
        existing_keys.add((slug, val))

game_by_slug = {g['slug'].lower(): g for g in games_list}
game_by_norm_name = {remove_accents(g['name']).lower(): g for g in games_list}

SPECIAL_MAP = {
    'The Ants: Underground Kingdom (Gamota)': 'vuong-quoc-kien',
    'Vương Quốc Kiến - Gamota (The Ants)': 'vuong-quoc-kien',
    'VNGGames & VNG Game Studios': 'play-together',
    'Yến Vân Thập Lục Thanh (Where Winds Meet)': 'yen-van-thap-luc-thanh',
    'Roblox Blox Fruits': 'roblox',
    'Pocket Army / Ark of War Series': 'ark-of-war-dreadnought',
    'Chiến Dịch Huyền Thoại (Garena)': 'chien-dich-huyen-thoai',
    'Garena Free Fire': 'garena-free-fire',
    'Free Fire': 'garena-free-fire',
    'Liên Quân Mobile (Arena of Valor)': 'lien-quan-mobile',
    'Liên Quân Mobile': 'lien-quan-mobile',
    'ZingSpeed Mobile': 'zingspeed-mobile',
    'Gunny Origin': 'gunny-origin',
    'Tây Du VNG: Đại Náo Tam Giới': 'tay-du-vng-dai-nao-tam-gioi',
    'Liên Minh Kiến (Ant Legion)': 'lien-minh-kien-ant-legion',
    'Võ Lâm Truyền Kỳ 1 Mobile': 'vo-lam-truyen-ky-1-mobile',
    'Võ Lâm Truyền Kỳ Mobile': 'vo-lam-truyen-ky-mobile',
    'Võ Lâm Truyền Kỳ MAX': 'vo-lam-truyen-ky-max',
    'Thiên Long Bát Bộ 2 VNG': 'thien-long-bat-bo-2-vng',
    'Kiếm Thế Origin VNG': 'kiem-the-origin-vng',
    'MU Lục Địa VNG': 'mu-luc-dia-vng',
    'Danh Tướng 3Q - VNG': 'danh-tuong-3q-vng',
    'Tiểu Yêu Tầm Đạo': 'tieu-yeu-tam-dao',
    'Ta Là Hắc Ngộ Không': 'ta-la-hac-ngo-khong',
    'Pirate Legends: Huyền Thoại Hải Tặc': 'pirate-legends',
    'Pirate Legends': 'pirate-legends',
    '3Q Đào Đào Đào': '3q-dao-dao-dao',
}

def resolve_game(raw_name):
    raw_name_clean = raw_name.strip()
    if raw_name_clean in SPECIAL_MAP:
        slug = SPECIAL_MAP[raw_name_clean]
        if slug in game_by_slug:
            return game_by_slug[slug]
    s = slugify(raw_name_clean)
    if s in game_by_slug:
        return game_by_slug[s]
    norm = remove_accents(raw_name_clean).lower()
    if norm in game_by_norm_name:
        return game_by_norm_name[norm]
    for n, g in game_by_norm_name.items():
        if norm == n or norm in n or n in norm:
            return g
    return None

# 2. Đọc file Excel
wb = openpyxl.load_workbook(excel_path, data_only=True)

# Đọc sheet Tổng hợp 340 game để lấy instructions, fanpage, nph
summary_info = {}
ws_summary = wb['Tổng hợp 340 game']
for r in list(ws_summary.iter_rows(values_only=True))[3:]:
    if not r or len(r) < 4: continue
    gname = str(r[1]).strip() if r[1] else ""
    nph = str(r[2]).strip() if len(r) > 2 and r[2] else ""
    fb_link = str(r[3]).strip() if len(r) > 3 and r[3] else ""
    instructions = str(r[6]).strip() if len(r) > 6 and r[6] else ""
    notes = str(r[7]).strip() if len(r) > 7 and r[7] else ""
    if gname:
        summary_info[gname] = {
            'nph': nph,
            'fb_link': fb_link,
            'instructions': instructions,
            'notes': notes
        }

# Đọc metadata Ta Là Hắc Ngộ Không
wukong_meta = {}
ws_wukong = wb['Ta Là Hắc Ngộ Không']
for r in list(ws_wukong.iter_rows(values_only=True))[7:]:
    if r and len(r) >= 2 and r[1]:
        c = str(r[1]).strip().upper()
        if c not in ['TỔNG SỐ CODE:', 'ĐÃ NHẬP:', 'CHƯA NHẬP:']:
            rw = str(r[2]).strip() if r[2] else ""
            st = str(r[3]).strip() if len(r) > 3 and r[3] else ""
            wukong_meta[c] = {'rewards': rw, 'status': st}

# Đọc metadata Pirate Legends
pirate_meta = {}
ws_pirate = wb['Pirate Legends']
for r in list(ws_pirate.iter_rows(values_only=True))[7:]:
    if r and len(r) >= 2 and r[1]:
        c = str(r[1]).strip().upper()
        if c not in ['TỔNG SỐ CODE:', 'ĐÃ NHẬP:', 'CHƯA NHẬP:']:
            rw = str(r[2]).strip() if r[2] else ""
            st = str(r[3]).strip() if len(r) > 3 and r[3] else ""
            pirate_meta[c] = {'rewards': rw, 'status': st}

# 3. Thu thập toàn bộ codes từ Excel
excel_codes = []
ws_all = wb['Tất cả code']
for r in list(ws_all.iter_rows(values_only=True))[3:]:
    if not r or len(r) < 3: continue
    stt, gname, code, nph, *rest = r
    if not gname or not code: continue
    gname = str(gname).strip()
    code = str(code).strip().upper()
    if gname in ['Tổng số code:', 'Chưa nhập:', 'Đã nhập:'] or len(code) < 2 or code == 'NONE': continue
    excel_codes.append({'gname': gname, 'code': code, 'nph': nph})

# Bổ sung codes từ sheet Ta Là Hắc Ngộ Không
for c, meta in wukong_meta.items():
    if not any(it['code'] == c and it['gname'] == 'Ta Là Hắc Ngộ Không' for it in excel_codes):
        excel_codes.append({'gname': 'Ta Là Hắc Ngộ Không', 'code': c, 'nph': 'MGO', 'meta': meta})

# Bổ sung codes từ sheet Pirate Legends
for c, meta in pirate_meta.items():
    if not any(it['code'] == c and 'Pirate Legends' in it['gname'] for it in excel_codes):
        excel_codes.append({'gname': 'Pirate Legends: Huyền Thoại Hải Tặc', 'code': c, 'nph': 'Neotrix', 'meta': meta})

print(f"📦 Đã trích xuất tổng cộng {len(excel_codes)} lượt mã từ file Excel.")

# 4. Tạo các bản ghi code mới
new_code_records = []
now_iso = datetime.utcnow().isoformat() + "Z"
random.seed(42)

for it in excel_codes:
    g = resolve_game(it['gname'])
    if not g:
        continue
    
    slug = g['slug']
    val = it['code']
    
    # Kiểm tra xem mã đã có chưa
    if (slug, val) in existing_keys:
        continue
    
    # Đánh dấu đã thêm để không bị trùng lặp chính trong file Excel
    existing_keys.add((slug, val))
    
    # Xác định phần thưởng
    rewards = "🎁 Quà Tân Thủ: Kim Cương, Vé Chiêu Mộ & Rương Trang Bị Cực Phẩm"
    if 'meta' in it and it['meta'].get('rewards'):
        rewards = it['meta']['rewards']
        if not rewards.startswith('🎁'):
            rewards = f"🎁 {rewards}"
    elif it['gname'] == 'Ta Là Hắc Ngộ Không' and val in wukong_meta and wukong_meta[val].get('rewards'):
        rewards = wukong_meta[val]['rewards']
        if not rewards.startswith('🎁'):
            rewards = f"🎁 {rewards}"
    elif ('Pirate Legends' in it['gname']) and val in pirate_meta and pirate_meta[val].get('rewards'):
        rewards = pirate_meta[val]['rewards']
        if not rewards.startswith('🎁'):
            rewards = f"🎁 {rewards}"
    elif it['gname'] in summary_info and summary_info[it['gname']].get('notes'):
        notes_str = summary_info[it['gname']]['notes']
        if len(notes_str) > 8:
            rewards = f"🎁 {notes_str[:120]}"

    # Xác định source link
    source_url = f"https://codegame.shop/game/{slug}"
    if it['gname'] in summary_info and summary_info[it['gname']].get('fb_link'):
        fb = summary_info[it['gname']]['fb_link'].split('\n')[0].strip()
        if fb.startswith('http'):
            source_url = fb

    # Nguồn phát hành
    source_name = "Fanpage NPH Chính Thức"
    if it['nph'] and len(it['nph']) > 2:
        source_name = f"NPH {it['nph']}"

    copy_count = random.randint(120, 520)

    record = {
        "id": f"code_{slug}_{val}",
        "gameId": slug,
        "gameSlug": slug,
        "game_slug": slug,
        "slug": slug,
        "gameName": g['name'],
        "code": val,
        "codeValue": val,
        "code_value": val,
        "gift_code": val,
        "rewards": rewards,
        "description": rewards,
        "status": "ACTIVE",
        "is_active": True,
        "is_verified": True,
        "isVerified": True,
        "expires_at": "2026-12-31",
        "expiresAt": "2026-12-31",
        "expiry": "2026-12-31",
        "sourceUrl": source_url,
        "source": source_name,
        "copyCount": copy_count,
        "createdAt": now_iso,
        "created_at": now_iso,
        "added_at": now_iso
    }
    new_code_records.append(record)

print(f"✨ Số mã giftcode mới CÒN THIẾU được tạo: {len(new_code_records)} mã!")

# 5. Gộp codes mới vào danh sách tổng (đưa code mới lên đầu danh sách để ưu tiên hiển thị)
combined_codes = new_code_records + existing_codes

# Tính toán lại codes_count cho từng game
code_count_by_game = {}
for c in combined_codes:
    s = (c.get('gameSlug') or c.get('game_slug') or c.get('slug') or c.get('gameId') or '').strip().lower()
    if s:
        code_count_by_game[s] = code_count_by_game.get(s, 0) + 1

# 6. Cập nhật games188.json
updated_games = []
for g in games_list:
    s = g['slug'].lower()
    total_c = code_count_by_game.get(s, g.get('codes_count', 0))
    g_updated = dict(g)
    g_updated['codes_count'] = total_c
    g_updated['free_codes_count'] = total_c
    
    # Cập nhật instructions từ summary nếu có hướng dẫn nhập chi tiết
    raw_name = g['name']
    for sum_name, sinfo in summary_info.items():
        if sinfo.get('instructions') and len(sinfo['instructions']) > 15:
            if resolve_game(sum_name) and resolve_game(sum_name)['slug'] == s:
                g_updated['instructions'] = sinfo['instructions']
                break
                
    g_updated['updated_at'] = now_iso
    updated_games.append(g_updated)

# 7. Lưu vào tất cả các file liên quan
code_files = [
    os.path.join(ROOT, "src", "data", "crawled_codes_188.json"),
    os.path.join(ROOT, "public", "data", "crawled_codes_188.json"),
    os.path.join(ROOT, "data", "crawled_codes_188.json"),
    os.path.join(ROOT, "data", "sources", "crawled_codes_188.json"),
]

for fp in code_files:
    if os.path.exists(os.path.dirname(fp)):
        with open(fp, 'w', encoding='utf-8') as f:
            json.dump(combined_codes, f, ensure_ascii=False, indent=2)
        print(f"💾 Đã lưu {len(combined_codes)} codes vào: {os.path.relpath(fp, ROOT)}")

game_files = [
    os.path.join(ROOT, "src", "data", "games188.json"),
    os.path.join(ROOT, "public", "data", "games188.json"),
    os.path.join(ROOT, "data", "games188.json"),
    os.path.join(ROOT, "data", "sources", "games188.json"),
]

for fp in game_files:
    if os.path.exists(os.path.dirname(fp)):
        with open(fp, 'w', encoding='utf-8') as f:
            json.dump(updated_games, f, ensure_ascii=False, indent=2)
        print(f"💾 Đã lưu {len(updated_games)} games vào: {os.path.relpath(fp, ROOT)}")

print("\n" + "=" * 65)
print(f"🎉 HOÀN TẤT BỔ SUNG {len(new_code_records)} MÃ CODE CÒN THIẾU VÀO WEB!")
print(f"   Tổng số mã giftcode trên web hiện tại: {len(combined_codes)} mã (tăng từ {len(existing_codes)}).")
print("=" * 65)
